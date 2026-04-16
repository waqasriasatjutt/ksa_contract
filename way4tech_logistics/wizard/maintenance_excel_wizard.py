import base64
import io
import logging

from odoo import api, fields, models, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    openpyxl = None
    _logger.warning(
        'openpyxl not found. Maintenance Excel import/export will not work. '
        'Install with: pip install openpyxl'
    )


MAINT_TYPE_MAP = {
    'preventive': 'Preventive',
    'corrective': 'Corrective',
    'accident': 'Accident Repair',
    'tyres': 'Tyres',
    'other': 'Other',
}
MAINT_TYPE_REVERSE = {v.lower(): k for k, v in MAINT_TYPE_MAP.items()}

STATE_MAP = {
    'draft': 'Draft',
    'confirmed': 'Confirmed',
    'done': 'Done',
}


class MaintenanceExcelExport(models.TransientModel):
    _name = 'way4tech.maintenance.excel.export'
    _description = 'Maintenance Excel Export Wizard'

    date_from = fields.Date(string='From')
    date_to = fields.Date(string='To')
    vehicle_id = fields.Many2one('fleet.vehicle', string='Vehicle (optional)')
    file_data = fields.Binary(string='Excel File', readonly=True)
    file_name = fields.Char(string='File Name', readonly=True)

    def action_export(self):
        self.ensure_one()
        if not openpyxl:
            raise UserError(_('openpyxl is required. Install with: pip install openpyxl'))

        domain = []
        if self.date_from:
            domain.append(('date', '>=', self.date_from))
        if self.date_to:
            domain.append(('date', '<=', self.date_to))
        if self.vehicle_id:
            domain.append(('vehicle_id', '=', self.vehicle_id.id))

        maintenances = self.env['fleet.vehicle.log.services'].search(domain, order='date desc')
        if not maintenances:
            raise UserError(_('No maintenance records found for the selected criteria.'))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Maintenance Records'

        headers = [
            'Reference', 'Date', 'Vehicle Plate', 'Vehicle Name',
            'Maintenance Category', 'Vendor', 'Amount',
            'Notes', 'State', 'Bill Reference', 'Analytic Account',
        ]

        header_fill = PatternFill(start_color='FFD700', end_color='FFD700', fill_type='solid')
        header_font = Font(bold=True)
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', wrap_text=True)

        row_idx = 2
        for m in maintenances:
            ws.cell(row=row_idx, column=1, value=m.way4tech_ref or '')
            ws.cell(row=row_idx, column=2, value=m.date)
            ws.cell(row=row_idx, column=3, value=m.vehicle_id.license_plate if m.vehicle_id else '')
            ws.cell(row=row_idx, column=4, value=m.vehicle_id.name if m.vehicle_id else '')
            ws.cell(row=row_idx, column=5, value=MAINT_TYPE_MAP.get(m.way4tech_maint_type, '') if m.way4tech_maint_type else '')
            ws.cell(row=row_idx, column=6, value=m.vendor_id.name if m.vendor_id else '')
            ws.cell(row=row_idx, column=7, value=float(m.amount or 0.0))
            ws.cell(row=row_idx, column=8, value=m.notes or '')
            ws.cell(row=row_idx, column=9, value=STATE_MAP.get(m.way4tech_state, '') if m.way4tech_state else 'Draft')
            ws.cell(row=row_idx, column=10, value=m.way4tech_bill_id.name if m.way4tech_bill_id else '')
            ws.cell(row=row_idx, column=11, value=m.way4tech_analytic_account_id.name if m.way4tech_analytic_account_id else '')
            row_idx += 1

        for col_idx in range(1, len(headers) + 1):
            column_letter = openpyxl.utils.get_column_letter(col_idx)
            ws.column_dimensions[column_letter].width = 20

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        file_content = output.read()
        output.close()

        self.write({
            'file_data': base64.b64encode(file_content),
            'file_name': 'maintenance_records.xlsx',
        })

        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/?model=way4tech.maintenance.excel.export&id=%d&field=file_data&filename_field=file_name&download=true' % self.id,
            'target': 'self',
        }


class MaintenanceExcelImport(models.TransientModel):
    _name = 'way4tech.maintenance.excel.import'
    _description = 'Maintenance Excel Import Wizard'

    excel_file = fields.Binary(string='Excel File (.xlsx)', required=True)
    file_name = fields.Char(string='File Name')
    import_log = fields.Text(string='Import Log', readonly=True)
    created_count = fields.Integer(string='Created', readonly=True)
    skipped_count = fields.Integer(string='Skipped', readonly=True)

    def action_import(self):
        self.ensure_one()
        if not openpyxl:
            raise UserError(_('openpyxl is required. Install with: pip install openpyxl'))
        if not self.excel_file:
            raise UserError(_('Please upload an Excel file.'))

        try:
            file_data = base64.b64decode(self.excel_file)
            wb = openpyxl.load_workbook(filename=io.BytesIO(file_data), read_only=True, data_only=True)
        except Exception as e:
            raise UserError(_('Could not open the Excel file: %s') % str(e))

        ws = wb.active
        created = 0
        skipped = 0
        errors = []

        # Expected columns (row 1 = header):
        # A=Date, B=Vehicle Plate, C=Maintenance Category, D=Vendor, E=Amount, F=Notes
        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(row):
                continue
            try:
                date_val = row[0] if len(row) > 0 else None
                plate = (row[1] or '').strip() if len(row) > 1 and row[1] else ''
                category_raw = (row[2] or '').strip() if len(row) > 2 and row[2] else ''
                vendor_name = (row[3] or '').strip() if len(row) > 3 and row[3] else ''
                amount = row[4] if len(row) > 4 else 0.0
                notes = (row[5] or '').strip() if len(row) > 5 and row[5] else ''

                if not plate:
                    errors.append('Row %d: License Plate is required — skipped.' % row_idx)
                    skipped += 1
                    continue
                if not date_val:
                    errors.append('Row %d: Date is required — skipped.' % row_idx)
                    skipped += 1
                    continue
                if not amount:
                    errors.append('Row %d: Amount is required — skipped.' % row_idx)
                    skipped += 1
                    continue

                vehicle = self.env['fleet.vehicle'].search(
                    [('license_plate', '=', plate)], limit=1
                )
                if not vehicle:
                    errors.append('Row %d: Vehicle with plate "%s" not found — skipped.' % (row_idx, plate))
                    skipped += 1
                    continue

                vendor = False
                if vendor_name:
                    vendor = self.env['res.partner'].search(
                        [('name', '=', vendor_name), ('supplier_rank', '>', 0)], limit=1
                    )
                    if not vendor:
                        # Auto-create vendor with supplier rank
                        vendor = self.env['res.partner'].create({
                            'name': vendor_name,
                            'supplier_rank': 1,
                        })

                maint_type = MAINT_TYPE_REVERSE.get(category_raw.lower()) or 'other'

                vals = {
                    'vehicle_id': vehicle.id,
                    'date': date_val,
                    'amount': float(amount or 0.0),
                    'way4tech_maint_type': maint_type,
                    'notes': notes,
                    'way4tech_state': 'draft',
                }
                if vendor:
                    vals['vendor_id'] = vendor.id

                self.env['fleet.vehicle.log.services'].create(vals)
                created += 1
            except Exception as e:
                errors.append('Row %d: %s' % (row_idx, str(e)))
                skipped += 1

        log_text = 'Created: %d\nSkipped: %d\n\n' % (created, skipped)
        if errors:
            log_text += 'Errors / warnings:\n' + '\n'.join(errors)

        self.write({
            'created_count': created,
            'skipped_count': skipped,
            'import_log': log_text,
        })

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'way4tech.maintenance.excel.import',
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
