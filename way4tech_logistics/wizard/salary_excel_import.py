import base64
import io
import logging

from odoo import api, fields, models, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    import openpyxl
except ImportError:
    openpyxl = None
    _logger.warning(
        'openpyxl library not found. Excel import functionality will not work. '
        'Install it with: pip install openpyxl'
    )

# ─────────────────────────────────────────────────────────────────────────────
# Expected Excel column order (row 1 = headers, row 2+ = data):
#
#  A  Employee Name          (required)
#  B  Platform               name — looked up in way4tech.platform.config
#  C  Valid Days             integer
#  D  Orders Completed       integer
#  E  Fixed Salary           float  (auto-computed if platform set; manual override ok)
#  F  Order Adjustment       float  (positive=bonus, negative=deduction; auto if platform)
#  G  Bonus                  float
#  H  Petrol Allowance       float
#  I  On-Time Deduction      float
#  J  Food Damage Deduction  float
#  K  Miss Day Penalty       float
#  L  Order Rejection Ded.   float
#  M  Misc Deduction         float
#  N  Advance                float
#  O  Fuel                   float
#  P  SIM Charges            float
#  Q  Loan                   float
#  R  Traffic Violation      float
#  S  Rent                   float
#  T  Note                   text
# ─────────────────────────────────────────────────────────────────────────────


class SalaryExcelImport(models.TransientModel):
    _name = 'way4tech.salary.excel.import'
    _description = 'Salary Excel Import Wizard'

    import_id = fields.Many2one(
        comodel_name='way4tech.salary.import',
        string='Salary Import',
        required=True,
    )
    excel_file = fields.Binary(
        string='Excel File (.xlsx)',
        required=True,
    )
    file_name = fields.Char(string='File Name')

    def action_import(self):
        self.ensure_one()

        if not openpyxl:
            raise UserError(_(
                'The openpyxl Python library is required for Excel import. '
                'Please install it: pip install openpyxl'
            ))
        if not self.excel_file:
            raise UserError(_('Please select an Excel file to import.'))

        try:
            file_data = base64.b64decode(self.excel_file)
        except Exception as e:
            raise UserError(_('Could not decode the file: %s') % str(e))

        try:
            workbook = openpyxl.load_workbook(
                filename=io.BytesIO(file_data), read_only=True, data_only=True
            )
        except Exception as e:
            raise UserError(_(
                'Could not open the Excel file. Please make sure it is a valid .xlsx file.\n'
                'Error: %s'
            ) % str(e))

        sheet = workbook.active

        def to_float(val):
            if val is None or val == '':
                return 0.0
            try:
                return float(val)
            except (ValueError, TypeError):
                return 0.0

        def to_int(val):
            if val is None or val == '':
                return 0
            try:
                return int(float(val))
            except (ValueError, TypeError):
                return 0

        def to_str(val):
            if val is None:
                return ''
            return str(val).strip()

        import_id = self.import_id
        default_platform = import_id.platform_id

        # Cache: platform name (lower) → way4tech.platform.config record
        platform_cache = {}

        def get_platform(name_raw):
            if not name_raw:
                return default_platform
            key = name_raw.strip().lower()
            if key not in platform_cache:
                rec = self.env['way4tech.platform.config'].search(
                    [('name', 'ilike', name_raw.strip())], limit=1
                )
                platform_cache[key] = rec
            return platform_cache[key]

        lines_to_create = []

        for row_index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(cell for cell in row):
                continue

            employee_name = to_str(row[0] if len(row) > 0 else None)
            if not employee_name:
                continue

            # Col B: platform
            platform_name = to_str(row[1] if len(row) > 1 else None)
            platform = get_platform(platform_name)

            # Col C-D: work data
            valid_days = to_int(row[2] if len(row) > 2 else None)
            orders_completed = to_int(row[3] if len(row) > 3 else None)

            # Col E-H: earnings (may be overridden by platform compute below)
            fixed_salary = to_float(row[4] if len(row) > 4 else None)
            order_adjustment = to_float(row[5] if len(row) > 5 else None)
            bonus = to_float(row[6] if len(row) > 6 else None)
            petrol_allowance = to_float(row[7] if len(row) > 7 else None)

            # Col I-M: performance deductions
            on_time_deduction = to_float(row[8] if len(row) > 8 else None)
            food_damage_deduction = to_float(row[9] if len(row) > 9 else None)
            miss_day_penalty = to_float(row[10] if len(row) > 10 else None)
            order_rejection_deduction = to_float(row[11] if len(row) > 11 else None)
            misc_deduction = to_float(row[12] if len(row) > 12 else None)

            # Col N-S: office deductions
            advance_deduction = to_float(row[13] if len(row) > 13 else None)
            fuel_deduction = to_float(row[14] if len(row) > 14 else None)
            sim_charges = to_float(row[15] if len(row) > 15 else None)
            loan = to_float(row[16] if len(row) > 16 else None)
            traffic_violation = to_float(row[17] if len(row) > 17 else None)
            rent = to_float(row[18] if len(row) > 18 else None)

            # Col T: note
            note = to_str(row[19] if len(row) > 19 else None)

            # Look up the employee's current outstanding loan balance for display
            loan_balance_before = 0.0

            # Auto-compute salary from platform rules if platform found
            # and fixed_salary was NOT explicitly provided in the file
            if platform and not fixed_salary and (valid_days or orders_completed):
                fixed_salary, order_adjustment = platform.compute_salary(
                    valid_days, orders_completed
                )

            # Try to find employee by name
            employee = self.env['hr.employee'].search(
                [('name', 'ilike', employee_name)], limit=1
            )

            line_vals = {
                'import_id': import_id.id,
                'sequence': row_index - 1,
                'employee_name': employee_name,
                'platform_id': platform.id if platform else False,
                'valid_days': valid_days,
                'orders_completed': orders_completed,
                'fixed_salary': fixed_salary,
                'order_adjustment': order_adjustment,
                'bonus': bonus,
                'petrol_allowance': petrol_allowance,
                'on_time_deduction': on_time_deduction,
                'food_damage_deduction': food_damage_deduction,
                'miss_day_penalty': miss_day_penalty,
                'order_rejection_deduction': order_rejection_deduction,
                'misc_deduction': misc_deduction,
                'advance_deduction': advance_deduction,
                'fuel_deduction': fuel_deduction,
                'sim_charges': sim_charges,
                'loan': loan,
                'loan_balance_before': loan_balance_before,
                'traffic_violation': traffic_violation,
                'rent': rent,
                'note': note,
                'has_error': False,
                'error_message': '',
            }

            if employee:
                line_vals['employee_id'] = employee.id
                # Look up outstanding loan balance for this employee
                loan_deductions = self.env['way4tech.employee.deduction'].search([
                    ('state', '=', 'posted'),
                    ('outstanding', '>', 0),
                    ('employee_id', '=', employee.id),
                    ('deduction_type', '=', 'loan'),
                ])
                line_vals['loan_balance_before'] = sum(
                    d.amount - d.recovered_amount for d in loan_deductions
                )
            else:
                # Auto-create a basic employee so the import can proceed without errors.
                # The payroll manager can fill in the contract and other details later.
                new_emp = self.env['hr.employee'].create({'name': employee_name})
                line_vals['employee_id'] = new_emp.id
                existing_note = line_vals.get('note') or ''
                line_vals['note'] = ('⚠ Auto-created employee — please review and add contract. ' + existing_note).rstrip()

            lines_to_create.append(line_vals)

        if not lines_to_create:
            raise UserError(_(
                'No data rows found in the Excel file. '
                'Make sure the file has data starting from row 2.'
            ))

        # Remove existing lines before importing new ones
        import_id.line_ids.unlink()

        # Create all lines
        self.env['way4tech.salary.import.line'].create(lines_to_create)

        import_id.state = 'imported'

        return {'type': 'ir.actions.act_window_close'}

    @api.model
    def action_download_template(self):
        """Generate and return a sample salary import XLSX template."""
        if not openpyxl:
            raise UserError(_(
                'The openpyxl Python library is required. '
                'Install it with: pip install openpyxl'
            ))

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Salary Import'

        from openpyxl.styles import Font, PatternFill, Alignment
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid')

        headers = [
            'Employee Name', 'Platform',
            'Valid Days', 'Orders Completed',
            'Basic Salary', 'Order Adjustment',
            'Bonus', 'Petrol Allowance',
            'On-Time Deduction', 'Food Damage Deduction', 'Miss Day Penalty',
            'Order Rejection Deduction', 'Misc Deduction',
            'Advance', 'Fuel', 'SIM Charges', 'Loan Installment',
            'Traffic Violation', 'Rent', 'Note',
        ]

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        # Sample data rows — mixed platforms in a single sheet
        sample_rows = [
            ['Ahmed Al-Rashidi',   'Hungerstation', 27, 310, '', '', 0, 0,   0, 0, 0, 0, 0,   0,   0,  0, 500, 0, 0, ''],
            ['Mohammed Al-Zahrani','Hungerstation', 25, 280, '', '', 0, 0,  10, 0, 0, 0, 0,   0,   0,  0,   0, 0, 0, ''],
            ['Khalid Al-Ghamdi',   'Hungerstation', 27, 300, '', '', 0, 0,   0, 0, 0, 0, 0,   0,   0,  0,   0, 0, 0, ''],
            ['Omar Al-Harthi',     'Keeta',         26, 290, '', '', 0, 0,   0, 5, 0, 0, 0, 200,   0, 50,   0, 0, 0, ''],
            ['Faisal Al-Qahtani',  'Keeta',         27, 320, '', '', 0, 0,   0, 0, 0, 0, 0,   0,   0,  0, 1000, 0, 0, 'Loan total 3000'],
            ['Saad Al-Dosari',     'Keeta',         27, 350, '', '', 0, 0,   0, 0, 0, 0, 0,   0, 100,  0,   0, 0, 0, ''],
        ]

        for row_data in sample_rows:
            ws.append(row_data)

        # Column widths
        col_widths = [22, 16, 11, 17, 13, 16, 7, 17, 16, 20, 16, 22, 15, 9, 7, 12, 17, 18, 7, 25]
        for i, width in enumerate(col_widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width

        # Write to bytes
        import io as _io
        output = _io.BytesIO()
        wb.save(output)
        output.seek(0)

        import base64 as _b64
        file_data = _b64.b64encode(output.read()).decode()

        # Return as file download attachment
        attachment = self.env['ir.attachment'].create({
            'name': 'salary_import_template.xlsx',
            'type': 'binary',
            'datas': file_data,
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })

        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%d?download=true' % attachment.id,
            'target': 'new',
        }
