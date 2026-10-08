import base64
import io
from datetime import date

from odoo import fields, models, _
from odoo.exceptions import UserError


class TruckProfitabilityWizard(models.TransientModel):
    """
    Wizard for generating Truck-wise and Customer-wise profitability PDF reports.
    The user selects a date range, grouping dimension (truck / client), and optional
    filters, then clicks Print Report to get a QWeb PDF.
    """
    _name = 'way4tech.truck.profitability.wizard'
    _description = 'Truck / Client Profitability Report'

    date_from = fields.Date(
        string='From Date',
        required=True,
        default=lambda self: fields.Date.today().replace(day=1),
    )
    date_to = fields.Date(
        string='To Date',
        required=True,
        default=fields.Date.today,
    )
    report_type = fields.Selection(
        selection=[
            ('truck', 'Truck-wise Profitability'),
            ('client', 'Customer-wise Profitability'),
        ],
        string='Group By',
        default='truck',
        required=True,
    )
    state_filter = fields.Selection(
        selection=[
            ('all', 'All Statuses'),
            ('confirmed', 'Confirmed Only'),
            ('done', 'Done Only'),
            ('confirmed_done', 'Confirmed & Done'),
        ],
        string='Trip Status',
        default='confirmed_done',
        required=True,
    )
    truck_ids = fields.Many2many(
        comodel_name='fleet.vehicle',
        string='Filter by Trucks',
        help='Leave empty to include all trucks.',
    )
    client_ids = fields.Many2many(
        comodel_name='res.partner',
        string='Filter by Clients',
        help='Leave empty to include all clients.',
    )
    xlsx_file = fields.Binary(string='Excel File', readonly=True, attachment=False)
    xlsx_filename = fields.Char(string='Excel File Name', readonly=True)
    include_maintenance = fields.Boolean(
        string='Include Maintenance Cost',
        default=True,
        help='Counts the maintenance logged against each truck in the period as '
             'a cost, so the profit is after maintenance. Every log is listed as '
             'its own line. Untick to report trip costs only.',
    )

    # ── helpers ───────────────────────────────────────────────────────────────

    def _build_domain(self):
        domain = [
            ('trip_date', '>=', self.date_from),
            ('trip_date', '<=', self.date_to),
            ('company_id', '=', self.env.company.id),
        ]
        if self.state_filter == 'confirmed':
            domain.append(('state', '=', 'confirmed'))
        elif self.state_filter == 'done':
            domain.append(('state', '=', 'done'))
        elif self.state_filter == 'confirmed_done':
            domain.append(('state', 'in', ['confirmed', 'done']))
        if self.truck_ids:
            domain.append(('truck_id', 'in', self.truck_ids.ids))
        if self.client_ids:
            domain.append(('client_id', 'in', self.client_ids.ids))
        return domain

    def _maintenance_domain(self):
        """Maintenance logged in the period, for the trucks in scope."""
        domain = [
            ('date', '>=', self.date_from),
            ('date', '<=', self.date_to),
            ('company_id', '=', self.env.company.id),
            ('amount', '!=', 0),
        ]
        # 2026-10-08: logs Odoo created by itself when a bill carrying a vehicle
        # was posted hold the price of the vehicle, not a maintenance cost.
        if 'way4tech_not_maintenance' in self.env[
                'fleet.vehicle.log.services']._fields:
            domain.append(('way4tech_not_maintenance', '=', False))
        if self.truck_ids:
            domain.append(('vehicle_id', 'in', self.truck_ids.ids))
        return domain

    def get_report_rows(self):
        """Every entry in the period as its own row, grouped by truck or client.

        One row per trip and one row per maintenance log, so ten maintenance
        logs read as ten lines. Profit on a row is its revenue less its driver,
        fuel, other and maintenance cost, and a group's totals are the sum of
        its rows.
        """
        self.ensure_one()
        trips = self.env['way4tech.truck.trip'].search(
            self._build_domain(), order='trip_date asc, id asc')
        logs = self.env['fleet.vehicle.log.services']
        if self.include_maintenance:
            logs = logs.search(self._maintenance_domain(), order='date asc, id asc')

        trip_types = dict(
            self.env['way4tech.truck.trip']._fields['trip_type'].selection or [])
        buckets = {}

        def bucket(key, label, sub_label):
            if key not in buckets:
                buckets[key] = {
                    'label': label or _('(Unassigned)'),
                    'sub_label': sub_label or '',
                    'rows': [],
                }
            return buckets[key]

        for trip in trips:
            if self.report_type == 'truck':
                key = ('truck', trip.truck_id.id)
                holder = bucket(key, trip.truck_id._way4tech_label(),
                                trip.truck_id.license_plate)
            else:
                key = ('client', trip.client_id.id)
                holder = bucket(key, trip.client_id.name,
                                trip.client_id.commercial_company_name)
            # Other Cost carries the trip's own other costs and, on a middleman
            # trip, the third-party hire. Left out of the profit, those costs
            # would simply disappear from the report.
            other = trip.other_cost_total - trip.fuel_cost
            holder['rows'].append({
                'date': trip.trip_date,
                'reference': trip.name or '',
                'truck': trip.truck_id._way4tech_label() if trip.truck_id else '',
                'client': trip.client_id.name or '',
                'kind': trip_types.get(trip.trip_type, trip.trip_type or ''),
                'revenue': trip.revenue,
                'driver': trip.driver_cost,
                'fuel': trip.fuel_cost,
                'other': other,
                'maintenance': 0.0,
            })

        for log in logs:
            if self.report_type == 'truck':
                key = ('truck', log.vehicle_id.id)
                holder = bucket(key, log.vehicle_id._way4tech_label(),
                                log.vehicle_id.license_plate)
            else:
                # Maintenance belongs to a truck, not to a customer, so it sits
                # in its own group and the grand total still ties.
                key = ('client', 0)
                holder = bucket(key, _('Vehicle Maintenance'), '')
            holder['rows'].append({
                'date': log.date,
                'reference': log.way4tech_ref or '',
                'truck': log.vehicle_id._way4tech_label() if log.vehicle_id else '',
                'client': '',
                'kind': _('Maintenance'),
                'revenue': 0.0,
                'driver': 0.0,
                'fuel': 0.0,
                'other': 0.0,
                'maintenance': log.amount,
            })

        groups = []
        for holder in buckets.values():
            rows = sorted(holder['rows'], key=lambda r: (r['date'] or date.min,
                                                         r['reference']))
            for row in rows:
                row['cost'] = row['driver'] + row['fuel'] + row['other'] + row['maintenance']
                row['profit'] = row['revenue'] - row['cost']
                row['margin'] = (row['profit'] / row['revenue'] * 100.0
                                 if row['revenue'] else 0.0)
            holder['rows'] = rows
            holder['totals'] = self._sum_rows(rows)
            groups.append(holder)
        groups.sort(key=lambda g: g['label'])
        return {
            'groups': groups,
            'totals': self._sum_rows([r for g in groups for r in g['rows']]),
        }

    def _sum_rows(self, rows):
        totals = {'count': len(rows)}
        for key in ('revenue', 'driver', 'fuel', 'other', 'maintenance',
                    'cost', 'profit'):
            totals[key] = sum(row.get(key, 0.0) for row in rows)
        totals['margin'] = (totals['profit'] / totals['revenue'] * 100.0
                            if totals['revenue'] else 0.0)
        return totals

    # ── actions ───────────────────────────────────────────────────────────────

    def action_export_xlsx(self):
        """The same report as a spreadsheet: a summary sheet and a detail sheet,
        the same columns, totals and order as the PDF."""
        self.ensure_one()
        if self.date_from > self.date_to:
            raise UserError(_('From Date must be on or before To Date.'))
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        except ImportError:
            raise UserError(_(
                'The openpyxl Python library is required for the Excel export.'))

        data = self.get_report_rows()
        symbol = self.env.company.currency_id.symbol or ''
        head_font = Font(bold=True, color='FFFFFF')
        head_fill = PatternFill('solid', fgColor='44546A')
        total_fill = PatternFill('solid', fgColor='FFF2CC')
        thin = Side(style='thin', color='B0B0B0')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        money = '#,##0.00'

        def put(sheet, row, values, bold=False, fill=None):
            for col, value in enumerate(values, 1):
                cell = sheet.cell(row=row, column=col, value=value)
                cell.border = border
                if isinstance(value, float):
                    cell.number_format = money
                    cell.alignment = Alignment(horizontal='right')
                else:
                    cell.alignment = Alignment(
                        horizontal='left', vertical='top', wrap_text=False)
                if bold:
                    cell.font = Font(bold=True)
                if fill:
                    cell.fill = fill
            return row + 1

        book = openpyxl.Workbook()
        summary = book.active
        summary.title = 'Summary'
        summary['A1'] = '%s  %s to %s' % (
            self._way4tech_report_title(), self.date_from, self.date_to)
        summary['A1'].font = Font(bold=True, size=13)
        headers = [self._way4tech_group_heading(), 'Entries',
                   'Revenue (%s)' % symbol, 'Driver Cost (%s)' % symbol,
                   'Fuel Cost (%s)' % symbol, 'Other Cost (%s)' % symbol,
                   'Maintenance (%s)' % symbol, 'Profit (%s)' % symbol,
                   'Margin%']
        row = 3
        for col, header in enumerate(headers, 1):
            cell = summary.cell(row=row, column=col, value=header)
            cell.font = head_font
            cell.fill = head_fill
            cell.border = border
        row += 1
        totals = data['totals']
        row = put(summary, row, [
            'PERIOD TOTAL', totals['count'], totals['revenue'], totals['driver'],
            totals['fuel'], totals['other'], totals['maintenance'],
            totals['profit'], round(totals['margin'], 1)], bold=True,
            fill=total_fill)
        for group in data['groups']:
            group_totals = group['totals']
            row = put(summary, row, [
                group['label'], group_totals['count'], group_totals['revenue'],
                group_totals['driver'], group_totals['fuel'],
                group_totals['other'], group_totals['maintenance'],
                group_totals['profit'], round(group_totals['margin'], 1)])
        for col, width in zip('ABCDEFGHI', (34, 9, 14, 14, 13, 13, 14, 14, 9)):
            summary.column_dimensions[col].width = width
        summary.freeze_panes = 'A4'

        detail = book.create_sheet('Detail')
        headers = ['Date', 'Reference', 'Truck', 'Client', 'Type',
                   'Revenue (%s)' % symbol, 'Driver Cost (%s)' % symbol,
                   'Fuel Cost (%s)' % symbol, 'Other Cost (%s)' % symbol,
                   'Maintenance (%s)' % symbol, 'Profit (%s)' % symbol,
                   'Margin%']
        row = 1
        for col, header in enumerate(headers, 1):
            cell = detail.cell(row=row, column=col, value=header)
            cell.font = head_font
            cell.fill = head_fill
            cell.border = border
        row += 1
        for group in data['groups']:
            row = put(detail, row, [group['label']], bold=True,
                      fill=PatternFill('solid', fgColor='E2EFDA'))
            for line in group['rows']:
                row = put(detail, row, [
                    line['date'] and str(line['date']) or '',
                    line['reference'], line['truck'], line['client'],
                    line['kind'], line['revenue'], line['driver'], line['fuel'],
                    line['other'], line['maintenance'], line['profit'],
                    round(line['margin'], 1) if line['revenue'] else ''])
            group_totals = group['totals']
            row = put(detail, row, [
                '', '', '', '', 'Subtotal', group_totals['revenue'],
                group_totals['driver'], group_totals['fuel'],
                group_totals['other'], group_totals['maintenance'],
                group_totals['profit'], round(group_totals['margin'], 1)],
                bold=True)
        for col, width in zip('ABCDEFGHIJKL',
                              (12, 24, 16, 30, 20, 14, 14, 13, 13, 14, 14, 9)):
            detail.column_dimensions[col].width = width
        detail.freeze_panes = 'A2'

        stream = io.BytesIO()
        book.save(stream)
        self.xlsx_file = base64.b64encode(stream.getvalue())
        self.xlsx_filename = '%s %s to %s.xlsx' % (
            self._way4tech_report_title(), self.date_from, self.date_to)
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s/%s/xlsx_file/%s?download=true' % (
                self._name, self.id, self.xlsx_filename),
            'target': 'self',
        }

    def _way4tech_report_title(self):
        return (_('Truck-wise Profitability Report')
                if self.report_type == 'truck'
                else _('Customer-wise Profitability Report'))

    def _way4tech_group_heading(self):
        return _('Truck') if self.report_type == 'truck' else _('Customer')

    def action_print_report(self):
        self.ensure_one()
        if self.date_from > self.date_to:
            raise UserError(_('From Date must be on or before To Date.'))
        if self.report_type == 'truck':
            ref = 'way4tech_logistics.action_report_truck_profitability'
        else:
            ref = 'way4tech_logistics.action_report_client_profitability'
        return self.env.ref(ref).report_action(self)
