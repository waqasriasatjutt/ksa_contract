"""
Run inside the Docker container to generate demo_salary_import.xlsx
  docker exec -it odoo19 python /mnt/extra-addons/way4tech_logistics/static/demo/gen_demo.py
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/mnt/extra-addons/way4tech_logistics/static/demo/demo_salary_import.xlsx"

wb = openpyxl.Workbook()

# ── Sheet 1: Salary Import Data ────────────────────────────────────────────────
ws = wb.active
ws.title = "Salary Import"

HEADERS = [
    "Employee Name",        # A
    "Rider Type",           # B  employee / freelancer
    "Platform",             # C  must match Platform Config name
    "Valid Days",           # D
    "Orders Completed",     # E
    "Fixed Salary",         # F  leave blank → auto-computed from platform
    "Order Adjustment",     # G  leave blank → auto-computed from platform
    "Bonus",                # H
    "Petrol Allowance",     # I
    "On-Time Deduction",    # J
    "Food Damage Deduction",# K
    "Miss Day Penalty",     # L
    "Order Rejection Ded.", # M
    "Misc Deduction",       # N
    "Advance Deduction",    # O
    "Fuel Deduction",       # P
    "SIM Charges",          # Q
    "Loan",                 # R
    "Traffic Violation",    # S
    "Rent",                 # T
    "Note",                 # U
]

# Colour fills
H_FILL  = PatternFill("solid", fgColor="1F3A5F")  # header: dark blue
EMP_F   = PatternFill("solid", fgColor="DBEAFE")  # employee info: light blue
EARN_F  = PatternFill("solid", fgColor="DCFCE7")  # earnings: light green
PERF_F  = PatternFill("solid", fgColor="FEF9C3")  # perf deductions: light yellow
OFF_F   = PatternFill("solid", fgColor="FEE2E2")  # office deductions: light red
NOTE_F  = PatternFill("solid", fgColor="F3F4F6")  # note: light grey
ALT_F   = PatternFill("solid", fgColor="F8FAFF")  # alternate row tint

COL_FILL = {
    1:EMP_F, 2:EMP_F, 3:EMP_F, 4:EMP_F, 5:EMP_F,
    6:EARN_F, 7:EARN_F, 8:EARN_F, 9:EARN_F,
    10:PERF_F, 11:PERF_F, 12:PERF_F, 13:PERF_F, 14:PERF_F,
    15:OFF_F, 16:OFF_F, 17:OFF_F, 18:OFF_F, 19:OFF_F, 20:OFF_F,
    21:NOTE_F,
}

H_FONT  = Font(color="FFFFFF", bold=True, size=10)
D_FONT  = Font(size=10)
B_FONT  = Font(size=10, bold=True)
C_ALIGN = Alignment(horizontal="center", vertical="center")
R_ALIGN = Alignment(horizontal="right")
L_ALIGN = Alignment(horizontal="left")

# Write headers
for c, h in enumerate(HEADERS, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill = H_FILL
    cell.font = H_FONT
    cell.alignment = C_ALIGN

# ── Demo rows ─────────────────────────────────────────────────────────────────
# Columns: Name, Type, Platform, Days, Orders, FixSal, OrdAdj, Bonus, Petrol,
#          OnTime, FoodDmg, MissDay, OrdRej, Misc,
#          Adv, Fuel, SIM, Loan, Traffic, Rent, Note
ROWS = [
    # ── CASE 1: Employee — platform auto-compute, meets min days, excess orders ──
    ["Ahmed Al-Rashidi",    "employee",   "Hungerstation", 28, 480,
     None, None,  100,  50,    0,   0,   0,   0,   0,    500,   0,   0,   0,   0,   0,
     "Full month — advance recovery 500"],

    # ── CASE 2: Employee — platform, below min-days → per-order rate ──────────
    ["Mohammed Al-Zahrani", "employee",   "Hungerstation", 22, 380,
     None, None,    0,  50,    0,  30,   0,   0,   0,      0,   0,  50,   0,   0,   0,
     "Below 27 days → per-order rate; food damage 30; SIM 50"],

    # ── CASE 3: Freelancer — platform auto-compute, excess orders ─────────────
    ["Fahad Al-Dosari",     "freelancer", "Keeta",         27, 510,
     None, None,    0,   0,   50,   0,   0,   0,   0,      0, 150,  50,   0,   0,   0,
     "Keeta freelancer — excess order bonus; fuel 150; SIM 50"],

    # ── CASE 4: Employee — manual salary, no platform ─────────────────────────
    ["Abdullah Al-Otaibi",  "employee",   "",              30,   0,
     1800,    0,  200, 100,    0,   0,   0,   0,   0,   1000,   0,   0, 500,   0,   0,
     "Manual fixed salary; advance 1000; loan 500"],

    # ── CASE 5: Employee — all deduction types ────────────────────────────────
    ["Khalid Al-Ghamdi",    "employee",   "Chefz",         27, 455,
     None, None,    0,  50,  100,  50,   0,   0,  25,    300, 200,  50,   0, 200, 800,
     "On-time 100; food 50; misc 25; adv 300; fuel 200; SIM 50; traffic 200; rent 800"],

    # ── CASE 6: Freelancer — manual, miss-day penalty ─────────────────────────
    ["Saleh Al-Harbi",      "freelancer", "",              26,   0,
     1500,    0,    0,   0,    0,   0, 150,   0,   0,      0,   0,  50,   0,   0,   0,
     "Freelancer manual salary; miss-day 150; SIM 50"],

    # ── CASE 7: Employee — top performer, high order bonus ────────────────────
    ["Omar Al-Shehri",      "employee",   "Hungerstation", 30, 620,
     None, None,  250, 100,    0,   0,   0,   0,   0,      0,   0,   0,   0,   0,   0,
     "620 orders vs 450 target — large order bonus + 250 bonus"],

    # ── CASE 8: Employee — food damage + order rejection + loan ───────────────
    ["Yusuf Al-Mutairi",    "employee",   "Keeta",         28, 430,
     None, None,    0,  50,    0, 200,   0, 150,   0,      0, 100,  50, 200,   0,   0,
     "Food damage 200; order rejection 150; fuel 100; loan 200"],

    # ── CASE 9: Freelancer — Keeta, SIM + fuel ────────────────────────────────
    ["Hassan Al-Anzi",      "freelancer", "Keeta",         29, 490,
     None, None,    0,   0,   50,   0,   0,   0,   0,      0, 300, 100,   0,   0,   0,
     "On-time 50; fuel 300; SIM 100"],

    # ── CASE 10: Employee — rent deduction ────────────────────────────────────
    ["Ibrahim Al-Subaie",   "employee",   "Chefz",         27, 440,
     None, None,    0,  50,  100,   0,   0,   0,   0,      0,   0,   0,   0,   0,1200,
     "On-time 100; rent 1200"],

    # ── CASE 11: Employee — short month ──────────────────────────────────────
    ["Nasser Al-Qahtani",   "employee",   "Hungerstation", 20, 290,
     None, None,    0,  50,    0,   0, 200,   0,   0,      0,   0,   0,   0,   0,   0,
     "Only 20 days worked; miss-day 200; per-order rate applied"],

    # ── CASE 12: Employee — misc + traffic ────────────────────────────────────
    ["Faisal Al-Shammari",  "employee",   "Hungerstation", 27, 450,
     None, None,    0,  50,    0,   0,   0,   0, 100,      0,   0,  50,   0, 300,   0,
     "Misc deduction 100; SIM 50; traffic violation 300"],

    # ── CASE 13: ERROR DEMO — name will not be found in Odoo ─────────────────
    ["New Rider Not In Odoo", "employee", "Keeta",         27, 460,
     None, None,    0,  50,    0,   0,   0,   0,   0,      0,   0,   0,   0,   0,   0,
     "DEMO ERROR: Name not matched — line will be RED; manually link employee"],

    # ── CASE 14: Freelancer — Chefz, full month ──────────────────────────────
    ["Saad Al-Bishi",       "freelancer", "Chefz",         28, 510,
     None, None,  150,   0,   50,   0,   0,   0,   0,      0, 200,   0,   0,   0,   0,
     "Chefz freelancer — bonus 150; on-time 50; fuel 200"],

    # ── CASE 15: Employee — order rejection + advance ─────────────────────────
    ["Waleed Al-Harthy",    "employee",   "Hungerstation", 27, 395,
     None, None,    0,  50,    0,   0,   0, 250,  50,    750,   0,  50,   0,   0,   0,
     "Order rejection 250; misc 50; advance recovery 750; SIM 50"],
]

# Write rows
thin = Side(style="thin", color="E5E7EB")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for r_idx, row in enumerate(ROWS, 2):
    for c_idx, val in enumerate(row, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=val)
        cell.font = D_FONT
        cell.border = border
        fill = COL_FILL.get(c_idx)
        if fill:
            cell.fill = fill
        # Alignment
        if c_idx in (4, 5):
            cell.alignment = C_ALIGN
        elif c_idx in range(6, 21):
            cell.alignment = R_ALIGN
        else:
            cell.alignment = L_ALIGN

# Column widths
WIDTHS = [24, 12, 16, 11, 16, 13, 14, 8, 8, 10, 14, 11, 14, 8, 14, 10, 10, 8, 14, 8, 46]
for i, w in enumerate(WIDTHS, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.row_dimensions[1].height = 36
ws.freeze_panes = "A2"

# ── Sheet 2: Legend & Instructions ────────────────────────────────────────────
ws2 = wb.create_sheet("Legend & Instructions")

LEGEND = [
    ("COLUMN GUIDE", "", "", ""),
    ("Column", "Field Name", "Required", "Notes"),
    ("A", "Employee Name",          "YES", "Must match hr.employee name (case-insensitive). Row turns RED if not found — link manually."),
    ("B", "Rider Type",             "No",  "'employee' or 'freelancer'. Defaults to employee if blank."),
    ("C", "Platform",               "No",  "Must match a Platform Configuration name. If set, Fixed Salary (F) and Order Adjustment (G) are auto-computed when F is blank."),
    ("D", "Valid Days",             "No",  "Working days in the month. Used in platform salary formula (threshold: min_days, default 27)."),
    ("E", "Orders Completed",       "No",  "Total orders delivered. Used to compute order bonus/penalty vs. min_orders target."),
    ("F", "Fixed Salary",           "No",  "Leave blank when Platform is set → auto-computed. Fill manually to override."),
    ("G", "Order Adjustment",       "No",  "Positive = bonus orders, Negative = short-order penalty. Auto-computed from Platform if F is blank."),
    ("H", "Bonus",                  "No",  "One-time bonus for this period."),
    ("I", "Petrol Allowance",       "No",  "Petrol/fuel allowance added to earnings."),
    ("J", "On-Time Deduction",      "No",  "Performance deduction for late deliveries."),
    ("K", "Food Damage Deduction",  "No",  "Deduction for damaged or spilled food orders."),
    ("L", "Miss Day Penalty",       "No",  "Penalty for absent working days."),
    ("M", "Order Rejection Ded.",   "No",  "Deduction for customer-rejected orders."),
    ("N", "Misc Deduction",         "No",  "Any other performance-related deduction."),
    ("O", "Advance Deduction",      "No",  "Recovery of previously issued advance (from Employee Advances & Deductions ledger)."),
    ("P", "Fuel Deduction",         "No",  "Recovery of fuel charges issued to employee."),
    ("Q", "SIM Charges",            "No",  "SIM card and mobile data charges."),
    ("R", "Loan",                   "No",  "Monthly loan repayment deduction."),
    ("S", "Traffic Violation",      "No",  "Traffic fine deducted from salary."),
    ("T", "Rent",                   "No",  "Accommodation rent deducted from salary."),
    ("U", "Note",                   "No",  "Internal note shown on the salary line (not printed on payslip)."),
    ("", "", "", ""),
    ("SALARY FORMULA", "", "", ""),
    ("Gross Earnings",   "= Fixed Salary + Order Adjustment + Bonus + Petrol Allowance", "", ""),
    ("Perf. Deductions", "= On-Time + Food Damage + Miss Day + Order Rejection + Misc",  "", ""),
    ("Office Deductions","= Advance + Fuel + SIM + Loan + Traffic + Rent",               "", ""),
    ("Net Payable",      "= Gross Earnings − Perf. Deductions − Office Deductions",      "", ""),
    ("", "", "", ""),
    ("COLOUR CODING (Salary Import sheet)", "", "", ""),
    ("Blue",    "Columns A–E: Employee identification",        "", ""),
    ("Green",   "Columns F–I: Earnings",                       "", ""),
    ("Yellow",  "Columns J–N: Performance deductions",         "", ""),
    ("Red/Pink","Columns O–T: Office / ledger deductions",     "", ""),
    ("Grey",    "Column U: Note",                              "", ""),
    ("", "", "", ""),
    ("PLATFORM AUTO-COMPUTE RULES", "", "", ""),
    ("IF valid_days >= min_days:",  "  basic_salary = fixed_salary",                               "", ""),
    ("",                            "  order_adj = (orders - min_orders) × excess_rate  [bonus]",  "", ""),
    ("",                            "           or (min_orders - orders) × short_rate   [penalty]","", ""),
    ("IF valid_days < min_days:",   "  basic_salary = orders_completed × per_order_rate",          "", ""),
    ("",                            "  order_adj = 0",                                             "", ""),
    ("", "", "", ""),
    ("DEMO DATA NOTES", "", "", ""),
    ("Row 14 (New Rider Not In Odoo)", "Name intentionally not in Odoo — demonstrates red error row.", "", ""),
    ("All other names",               "Replace with actual hr.employee names in your Odoo database.", "", ""),
    ("Platform names (C column)",     "Replace with your actual Platform Configuration names.",       "", ""),
]

L_H_FILL  = PatternFill("solid", fgColor="1F3A5F")
L_S_FILL  = PatternFill("solid", fgColor="374151")
L_H_FONT  = Font(color="FFFFFF", bold=True, size=10)
L_S_FONT  = Font(color="FFFFFF", bold=True, size=10)
L_D_FONT  = Font(size=10)
L_B_FONT  = Font(size=10, bold=True)

for r_idx, (col_a, col_b, col_c, col_d) in enumerate(LEGEND, 1):
    is_section = col_a in ("COLUMN GUIDE", "SALARY FORMULA", "COLOUR CODING (Salary Import sheet)",
                           "PLATFORM AUTO-COMPUTE RULES", "DEMO DATA NOTES")
    is_header  = col_b == "Field Name"

    for c_idx, val in enumerate([col_a, col_b, col_c, col_d], 1):
        cell = ws2.cell(row=r_idx, column=c_idx, value=val)
        if is_section:
            cell.fill = L_S_FILL
            cell.font = L_S_FONT
        elif is_header:
            cell.fill = L_H_FILL
            cell.font = L_H_FONT
        else:
            cell.font = L_D_FONT if c_idx != 1 else L_B_FONT
        cell.alignment = L_ALIGN

ws2.column_dimensions["A"].width = 34
ws2.column_dimensions["B"].width = 64
ws2.column_dimensions["C"].width = 10
ws2.column_dimensions["D"].width = 72
ws2.freeze_panes = "A3"

# ── Sheet 3: Sample Platform Config reference ──────────────────────────────────
ws3 = wb.create_sheet("Platform Reference")

ws3.cell(1, 1, "PLATFORM CONFIGURATION REFERENCE (sample — create matching records in Odoo)").font = Font(bold=True, size=11, color="1F3A5F")

PLAT_HEADERS = ["Platform Name", "Min Days", "Fixed Salary (SAR)", "Per-Order Rate (SAR)",
                 "Min Orders", "Short Order Rate", "Excess Order Rate", "Notes"]
for c, h in enumerate(PLAT_HEADERS, 1):
    cell = ws3.cell(2, c, h)
    cell.fill = H_FILL
    cell.font = H_FONT
    cell.alignment = C_ALIGN

PLAT_DATA = [
    ["Hungerstation", 27, 1800, 4.0, 450, 1.0, 1.5, "Most common platform in KSA"],
    ["Keeta",         27, 1700, 3.8, 450, 0.8, 1.2, ""],
    ["Chefz",         27, 1600, 3.5, 400, 1.0, 1.0, "Lower target 400 orders"],
]
for r_idx, row in enumerate(PLAT_DATA, 3):
    for c_idx, val in enumerate(row, 1):
        cell = ws3.cell(r_idx, c_idx, val)
        cell.font = Font(size=10)
        cell.alignment = R_ALIGN if isinstance(val, (int, float)) else L_ALIGN

for c, w in enumerate([18, 10, 20, 18, 12, 17, 17, 30], 1):
    ws3.column_dimensions[get_column_letter(c)].width = w

ws3.cell(7, 1, "⚠  Create these platforms in: Way4Tech Logistics → Configuration → Platform Configurations").font = Font(bold=True, color="C0392B", size=10)
ws3.cell(8, 1, "⚠  The 'Platform' column (C) in the Salary Import sheet must match these names exactly.").font = Font(bold=True, color="C0392B", size=10)

# Save
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(f"✅  Saved: {OUT}")
print(f"    Rows: {len(ROWS)} demo salary lines across 3 sheets")
